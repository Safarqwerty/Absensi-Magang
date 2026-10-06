from flask import Flask, render_template, request, redirect, url_for, session, send_from_directory, send_file, flash
from werkzeug.utils import secure_filename
from datetime import datetime, timedelta, date, timezone
import os
import re
import pymysql
import base64
import math
from io import BytesIO
from PIL import Image as PILImage

# IMPORT UNTUK EXCEL (OPENPYXL)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage


# ---------------- ZONA WAKTU: WITA (UTC+8) ----------------
# Semua pencatatan waktu memakai WITA, tidak bergantung pada zona waktu server.
try:
    from zoneinfo import ZoneInfo
    WITA = ZoneInfo("Asia/Makassar")
except Exception:                       # Windows tanpa paket tzdata, atau Python < 3.9
    WITA = timezone(timedelta(hours=8))


def sekarang():
    """Waktu saat ini di WITA (tanpa info zona, agar cocok dengan format yang disimpan di database)."""
    return datetime.now(WITA).replace(tzinfo=None)


app = Flask(__name__)
app.secret_key = "ksop_makassar_secure_key_2026"

UPLOAD_FOLDER = 'uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

KANTOR_LAT = -5.115395      # pin "Kantor KSOP Utama Makassar" (sama dengan peta di halaman user)
KANTOR_LON = 119.413559
MAX_RADIUS_METER = 300

BATAS_MASUK = "07:30:00"    # absen masuk lewat jam ini dicatat Telat
BATAS_PULANG = "16:00:00"   # absen pulang baru bisa dilakukan mulai jam ini


def hitung_jarak_meter(lat1, lon1, lat2, lon2):
    """Jarak dua titik koordinat dalam meter (rumus haversine)."""
    R = 6371000
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return R * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def get_db_connection():
    return pymysql.connect(
        user='manj1947_admin',
        password='W)33T$5]9H~EV+)L',
        database='manj1947_absensi-ksop',
        cursorclass=pymysql.cursors.DictCursor
    )


def init_db():
    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INT AUTO_INCREMENT PRIMARY KEY,
                username VARCHAR(100) UNIQUE,
                password VARCHAR(100),
                nama VARCHAR(255),
                posisi VARCHAR(100),
                asal_instansi VARCHAR(255),
                mentor VARCHAR(255),
                email VARCHAR(255),
                role VARCHAR(50),
                foto_wajah VARCHAR(255),
                kategori_magang VARCHAR(100) DEFAULT 'Magang Lainnya'
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS riwayat (
                id INT AUTO_INCREMENT PRIMARY KEY,
                nama VARCHAR(255),
                posisi VARCHAR(100),
                status VARCHAR(100),
                alasan TEXT,
                bukti VARCHAR(255),
                waktu VARCHAR(100),
                latitude VARCHAR(50),
                longitude VARCHAR(50),
                metode_verifikasi VARCHAR(255),
                status_persetujuan VARCHAR(50) DEFAULT 'Menunggu',
                catatan_admin TEXT,
                kategori_magang VARCHAR(100) DEFAULT 'Magang Lainnya'
            )
        """)

        for col, dtype in [('kategori_magang', "VARCHAR(100) DEFAULT 'Magang Lainnya'"),
                           ('status_persetujuan', "VARCHAR(50) DEFAULT 'Menunggu'"),
                           ('catatan_admin', 'TEXT')]:
            try:
                cursor.execute(f"ALTER TABLE riwayat ADD COLUMN {col} {dtype}")
                conn.commit()
            except Exception:
                pass

        cursor.execute("SELECT * FROM users WHERE username = 'admin'")
        if not cursor.fetchone():
            cursor.execute(
                "INSERT INTO users (username, password, nama, posisi, asal_instansi, mentor, email, role, foto_wajah, kategori_magang) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
                ('admin', 'admin123', 'Administrator KSOP', 'Pusat Data & Informasi', 'KSOP Utama Makassar', 'Kepala PUSDATIN', 'admin.ksop@dephub.go.id', 'admin', '-', 'Admin Pusat')
            )
            conn.commit()

        conn.close()
    except Exception as e:
        print(f"Database Error: {e}")


init_db()


def tambah_filter_waktu(query, params, filter_tanggal, filter_bulan, filter_tahun):
    """Menambahkan filter tanggal / bulan / tahun ke query.
    Jika hanya bulan yang dipilih, tahun berjalan dipakai otomatis."""
    if filter_tanggal:
        query += " AND DATE(waktu) = %s"
        params.append(filter_tanggal)
    elif filter_bulan:
        tahun = filter_tahun or str(sekarang().year)
        query += " AND MONTH(waktu) = %s AND YEAR(waktu) = %s"
        params.extend([filter_bulan, tahun])
    elif filter_tahun:
        query += " AND YEAR(waktu) = %s"
        params.append(filter_tahun)
    return query, params


def kembali_aman(nilai):
    """Hanya izinkan redirect ke path internal."""
    if not nilai or not nilai.startswith("/") or nilai.startswith("//"):
        return url_for('index')
    return nilai


def hapus_file_upload(nama_file):
    """Hapus file di folder uploads dengan aman (tanpa path traversal)."""
    if not nama_file or nama_file == '-':
        return
    aman = os.path.basename(nama_file)
    path = os.path.join(app.config['UPLOAD_FOLDER'], aman)
    try:
        if os.path.isfile(path):
            os.remove(path)
    except Exception:
        pass


# ---------------- REKAP ABSENSI & LENCANA ----------------
NAMA_BULAN = ['Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni', 'Juli',
              'Agustus', 'September', 'Oktober', 'November', 'Desember']
LENCANA = [1, 2, 3]   # peringkat; ikon medali digambar di template (SVG)
STATUS_MASUK = ('Masuk', 'Telat')            # dipakai untuk menentukan jam masuk
STATUS_HADIR = ('Masuk', 'Telat', 'Hadir')   # dihitung sebagai hari hadir
EXT_GAMBAR = ('.jpg', '.jpeg', '.png', '.gif', '.webp', '.bmp')


def detik_dari_waktu(w):
    try:
        t = datetime.strptime(str(w)[:19], "%Y-%m-%d %H:%M:%S")
        return t.hour * 3600 + t.minute * 60 + t.second
    except Exception:
        return None


def format_detik(d):
    d = int(d)
    return f"{d // 3600:02d}:{(d % 3600) // 60:02d}:{d % 60:02d}"


def format_tanggal(d):
    return f"{d.day:02d} {NAMA_BULAN[d.month - 1][:3]} {d.year}"


def label_periode(dari, sampai):
    if dari == sampai:
        return format_tanggal(dari)
    return f"{format_tanggal(dari)} - {format_tanggal(sampai)}"


def baca_rentang_rekap(arg_dari, arg_sampai):
    """Baca rentang tanggal dari URL. Default: tanggal 1 bulan ini s/d hari ini."""
    hari_ini = sekarang().date()
    try:
        dari = datetime.strptime(arg_dari, "%Y-%m-%d").date()
    except Exception:
        dari = hari_ini.replace(day=1)
    try:
        sampai = datetime.strptime(arg_sampai, "%Y-%m-%d").date()
    except Exception:
        sampai = hari_ini
    if dari > sampai:
        dari, sampai = sampai, dari
    return dari, sampai


def preset_rekap():
    h = sekarang().date()
    awal_bulan = h.replace(day=1)
    akhir_bln_lalu = awal_bulan - timedelta(days=1)
    return [
        ("Hari Ini", h, h),
        ("7 Hari Terakhir", h - timedelta(days=6), h),
        ("Bulan Ini", awal_bulan, h),
        ("Bulan Lalu", akhir_bln_lalu.replace(day=1), akhir_bln_lalu),
        ("30 Hari Terakhir", h - timedelta(days=29), h),
        ("Tahun Ini", h.replace(month=1, day=1), h),
    ]


def tentukan_rentang(tanggal, bulan, tahun, arg_dari, arg_sampai):
    """Filter tanggal/bulan/tahun (kalau diisi) menang atas rentang manual."""
    if tanggal:
        try:
            d = datetime.strptime(tanggal, "%Y-%m-%d").date()
            return d, d
        except Exception:
            pass
    if bulan or tahun:
        th = int(tahun) if tahun and tahun.isdigit() else sekarang().year
        if bulan and bulan.isdigit() and 1 <= int(bulan) <= 12:
            b = int(bulan)
            awal = date(th, b, 1)
            akhir = (date(th + 1, 1, 1) if b == 12 else date(th, b + 1, 1)) - timedelta(days=1)
            return awal, akhir
        return date(th, 1, 1), date(th, 12, 31)
    return baca_rentang_rekap(arg_dari, arg_sampai)


def ambil_daftar_rekap(dari, sampai, kategori='', cari=''):
    """Semua baris absensi (termasuk bukti/foto) sesuai filter rekap."""
    conn = get_db_connection()
    cursor = conn.cursor()
    query = "SELECT * FROM riwayat WHERE DATE(waktu) BETWEEN %s AND %s"
    params = [dari.strftime("%Y-%m-%d"), sampai.strftime("%Y-%m-%d")]
    if kategori:
        query += " AND kategori_magang = %s"
        params.append(kategori)
    if cari:
        query += " AND nama LIKE %s"
        params.append(f"%{cari}%")
    query += " ORDER BY waktu DESC"
    cursor.execute(query, params)
    data = cursor.fetchall()
    conn.close()
    return data


def tulis_sheet_daftar(wb, data, judul="Daftar Absensi"):
    """Sheet berisi daftar absensi + foto tertanam + link bukti."""
    ws = wb.create_sheet(judul)
    if not hasattr(wb, '_buf'):
        wb._buf = []          # jaga buffer gambar tetap hidup sampai wb.save()

    header = ["No", "Tanggal", "Jam", "Nama", "Kategori Magang", "Posisi", "Status",
              "Alasan", "Persetujuan", "Catatan Admin", "Foto / Bukti", "Link Bukti"]
    ws.append(header)
    for col in range(1, len(header) + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="0B2545", end_color="0B2545", fill_type="solid")
        c.alignment = Alignment(horizontal="center")

    if not data:
        ws.append(["", "Tidak ada data pada filter yang dipilih"])
    for i, row in enumerate(data, start=1):
        waktu = str(row['waktu'])
        ws.append([i, waktu[:10], waktu[11:19], row['nama'],
                   row.get('kategori_magang') or "Magang Lainnya", row['posisi'], row['status'],
                   row['alasan'] or "-", row['status_persetujuan'] or "Menunggu",
                   row['catatan_admin'] or "-"])
        r = ws.max_row
        for c in range(1, 13):
            ws.cell(row=r, column=c).alignment = Alignment(vertical="center", wrap_text=True)

        bukti = row.get('bukti')
        if bukti and bukti != '-':
            path = os.path.join(app.config['UPLOAD_FOLDER'], os.path.basename(bukti))
            link = ws.cell(row=r, column=12, value="Buka")
            link.hyperlink = url_for('uploaded_file', filename=bukti, _external=True)
            link.font = Font(color="0096C7", underline="single")
            if bukti.lower().endswith(EXT_GAMBAR) and os.path.isfile(path):
                try:
                    im = PILImage.open(path).convert('RGB')
                    im.thumbnail((140, 140))
                    buf = BytesIO()
                    buf.name = "foto.jpg"
                    im.save(buf, 'JPEG', quality=80)
                    buf.seek(0)
                    wb._buf.append(buf)
                    xl = XLImage(buf)
                    xl.width, xl.height = im.size
                    ws.add_image(xl, f"K{r}")
                    ws.row_dimensions[r].height = im.size[1] * 0.75 + 6
                except Exception:
                    ws.cell(row=r, column=11, value="(foto gagal dimuat)")
            else:
                ws.cell(row=r, column=11, value="Berkas: " + bukti)
        else:
            ws.cell(row=r, column=11, value="-")

    for i, w in enumerate([5, 12, 10, 22, 18, 20, 12, 28, 14, 24, 24, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def top_tercepat_hari_ini(kategori='', cari=''):
    """3 orang dengan absen masuk paling awal hari ini (absen yang ditolak tidak dihitung)."""
    hari_ini = sekarang().strftime("%Y-%m-%d")
    conn = get_db_connection()
    cursor = conn.cursor()
    query = ("SELECT nama, kategori_magang, status, waktu FROM riwayat "
             "WHERE waktu LIKE %s AND status IN ('Masuk', 'Telat') "
             "AND (status_persetujuan IS NULL OR status_persetujuan != 'Ditolak')")
    params = [hari_ini + '%']
    if kategori:
        query += " AND kategori_magang = %s"
        params.append(kategori)
    if cari:
        query += " AND nama LIKE %s"
        params.append(f"%{cari}%")
    query += " ORDER BY waktu ASC"
    cursor.execute(query, params)
    rows = cursor.fetchall()
    conn.close()

    hasil, sudah = [], set()
    for r in rows:
        if r['nama'] in sudah:
            continue
        sudah.add(r['nama'])
        hasil.append({'nama': r['nama'], 'kategori': r['kategori_magang'] or 'Magang Lainnya',
                      'jam': str(r['waktu'])[11:19], 'status': r['status'],
                      'lencana': LENCANA[len(hasil)]})
        if len(hasil) == 3:
            break
    return hasil


def ringkasan_hari_ini(kategori='', cari=''):
    """Hitungan kartu dashboard admin untuk HARI INI (WITA). Otomatis mulai dari 0 setiap hari baru.
    Absen yang Ditolak tidak dihitung. Mengikuti filter kategori & nama, tetapi tidak filter tanggal."""
    hari_ini = sekarang().strftime("%Y-%m-%d")
    conn = get_db_connection()
    cursor = conn.cursor()
    query = ("SELECT status, COUNT(*) AS jml FROM riwayat "
             "WHERE waktu LIKE %s AND (status_persetujuan IS NULL OR status_persetujuan != 'Ditolak')")
    params = [hari_ini + '%']
    if kategori:
        query += " AND kategori_magang = %s"
        params.append(kategori)
    if cari:
        query += " AND nama LIKE %s"
        params.append(f"%{cari}%")
    query += " GROUP BY status"
    cursor.execute(query, params)
    hitung = {r['status']: r['jml'] for r in cursor.fetchall()}
    conn.close()

    return {
        'tanggal': format_tanggal(sekarang().date()),
        'total_log': sum(hitung.values()),
        'hadir': hitung.get('Masuk', 0) + hitung.get('Hadir', 0),   # masuk tepat waktu
        'telat': hitung.get('Telat', 0),
        'pulang': hitung.get('Pulang', 0),
        'izin_sakit': hitung.get('Izin', 0) + hitung.get('Sakit', 0),
        'cuti': hitung.get('Cuti', 0),
        'tanpa_ket': hitung.get('Tanpa Keterangan', 0),
    }


def hitung_rekap_periode(dari, sampai, kategori='', cari=''):
    """Rekap per peserta untuk rentang tanggal [dari, sampai], termasuk rincian tanggal & jam per hari,
    + 3 peserta dengan rata-rata jam masuk tercepat. Mengikuti filter kategori dan nama."""
    conn = get_db_connection()
    cursor = conn.cursor()
    query = ("SELECT nama, kategori_magang, status, waktu FROM riwayat "
             "WHERE DATE(waktu) BETWEEN %s AND %s "
             "AND (status_persetujuan IS NULL OR status_persetujuan != 'Ditolak')")
    params = [dari.strftime("%Y-%m-%d"), sampai.strftime("%Y-%m-%d")]
    if kategori:
        query += " AND kategori_magang = %s"
        params.append(kategori)
    if cari:
        query += " AND nama LIKE %s"
        params.append(f"%{cari}%")
    cursor.execute(query, params)
    rows = cursor.fetchall()

    uq = "SELECT nama, kategori_magang FROM users WHERE role != 'admin'"
    up = []
    if kategori:
        uq += " AND kategori_magang = %s"
        up.append(kategori)
    if cari:
        uq += " AND nama LIKE %s"
        up.append(f"%{cari}%")
    cursor.execute(uq + " ORDER BY nama", up)
    users = cursor.fetchall()
    conn.close()

    def kosong(nama, kat):
        return {'nama': nama, 'kategori': kat or 'Magang Lainnya',
                'hari_hadir': set(), 'hari_telat': set(), 'jam_masuk': {}, 'harian': {},
                'izin': 0, 'sakit': 0, 'cuti': 0, 'tanpa_ket': 0}

    data = {u['nama']: kosong(u['nama'], u['kategori_magang']) for u in users}

    for r in rows:
        d = data.setdefault(r['nama'], kosong(r['nama'], r['kategori_magang']))
        tanggal = str(r['waktu'])[:10]
        jam_str = str(r['waktu'])[11:19]
        st = r['status']

        # rincian per hari: tanggal + jam masuk + jam pulang
        h = d['harian'].setdefault(tanggal, {'masuk': '-', 'pulang': '-', 'ket': ''})
        if st in STATUS_MASUK:
            if h['masuk'] == '-' or jam_str < h['masuk']:
                h['masuk'] = jam_str
                h['ket'] = st
        elif st == 'Pulang':
            if h['pulang'] == '-' or jam_str > h['pulang']:
                h['pulang'] = jam_str
        else:
            if not h['ket']:
                h['ket'] = st
            if h['masuk'] == '-':
                h['masuk'] = jam_str

        if st in STATUS_HADIR:
            d['hari_hadir'].add(tanggal)
        if st == 'Telat':
            d['hari_telat'].add(tanggal)
        if st in STATUS_MASUK:
            detik = detik_dari_waktu(r['waktu'])
            if detik is not None and (tanggal not in d['jam_masuk'] or detik < d['jam_masuk'][tanggal]):
                d['jam_masuk'][tanggal] = detik
        elif st == 'Izin':
            d['izin'] += 1
        elif st == 'Sakit':
            d['sakit'] += 1
        elif st == 'Cuti':
            d['cuti'] += 1
        elif st == 'Tanpa Keterangan':
            d['tanpa_ket'] += 1

    rekap = []
    for d in data.values():
        rata = (sum(d['jam_masuk'].values()) / len(d['jam_masuk'])) if d['jam_masuk'] else None
        rincian = []
        for t, v in sorted(d['harian'].items()):
            try:
                tgl_txt = format_tanggal(datetime.strptime(t, "%Y-%m-%d").date())
            except Exception:
                tgl_txt = t
            rincian.append({'tanggal': tgl_txt, 'masuk': v['masuk'], 'pulang': v['pulang'],
                            'ket': v['ket'] or '-'})
        rekap.append({
            'nama': d['nama'], 'kategori': d['kategori'],
            'hadir': len(d['hari_hadir']), 'telat': len(d['hari_telat']),
            'izin': d['izin'], 'sakit': d['sakit'], 'cuti': d['cuti'], 'tanpa_ket': d['tanpa_ket'],
            'rata_detik': rata,
            'rata_jam': format_detik(rata) if rata is not None else '-',
            'rincian': rincian,
            'lencana': 0
        })

    # 3 tercepat: rata-rata jam masuk paling awal (seri -> hari hadir lebih banyak)
    kandidat = sorted([r for r in rekap if r['rata_detik'] is not None],
                      key=lambda r: (r['rata_detik'], -r['hadir']))[:3]
    for i, r in enumerate(kandidat):
        r['lencana'] = LENCANA[i]

    rekap.sort(key=lambda r: (-r['hadir'], r['nama'].lower()))
    return rekap, kandidat


def render_user_dashboard(error=None, success=None):
    """Render dashboard user dengan SEMUA variabel yang dibutuhkan template,
    sehingga pesan error/sukses bisa ditampilkan tanpa merusak tampilan."""
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM riwayat WHERE nama = %s ORDER BY id DESC", (session['nama'],))
    riwayat_user = cursor.fetchall()
    conn.close()

    hari_ini = sekarang().strftime("%Y-%m-%d")
    absen_masuk_hari_ini = None
    absen_pulang_hari_ini = None

    for r in riwayat_user:
        if hari_ini in str(r['waktu']):
            if r['status'] in ['Masuk', 'Telat'] and not absen_masuk_hari_ini:
                absen_masuk_hari_ini = r
            elif r['status'] == 'Pulang' and not absen_pulang_hari_ini:
                absen_pulang_hari_ini = r

    total_hadir = sum(1 for r in riwayat_user if r['status'] in ['Masuk', 'Pulang', 'Hadir', 'Telat'])
    total_izin = sum(1 for r in riwayat_user if r['status'] == 'Izin')
    total_sakit = sum(1 for r in riwayat_user if r['status'] == 'Sakit')
    total_cuti = sum(1 for r in riwayat_user if r['status'] == 'Cuti')
    total_tanpa_ket = sum(1 for r in riwayat_user if r['status'] == 'Tanpa Keterangan')

    return render_template("user_dashboard.html",
                           user=session,
                           total_hadir=total_hadir,
                           total_izin=total_izin,
                           total_sakit=total_sakit,
                           total_cuti=total_cuti,
                           total_tanpa_ket=total_tanpa_ket,
                           absen_masuk=absen_masuk_hari_ini,
                           absen_pulang=absen_pulang_hari_ini,
                           error=error,
                           success=success,
                           kantor_lat=KANTOR_LAT,
                           kantor_lon=KANTOR_LON,
                           max_radius=MAX_RADIUS_METER,
                           batas_masuk=BATAS_MASUK[:5],
                           batas_pulang=BATAS_PULANG[:5])


@app.route("/")
def index():
    if 'username' not in session:
        return redirect(url_for('login'))

    if session.get('role') == 'admin':
        filter_tanggal = request.args.get('tanggal', '')
        filter_bulan = request.args.get('bulan', '')
        filter_tahun = request.args.get('tahun', '')
        filter_kategori = request.args.get('kategori', '')
        filter_cari = request.args.get('cari', '').strip()

        conn = get_db_connection()
        cursor = conn.cursor()

        query = "SELECT * FROM riwayat WHERE 1=1"
        params = []
        query, params = tambah_filter_waktu(query, params, filter_tanggal, filter_bulan, filter_tahun)

        if filter_kategori:
            query += " AND kategori_magang = %s"
            params.append(filter_kategori)

        if filter_cari:
            query += " AND nama LIKE %s"
            params.append(f"%{filter_cari}%")

        query += " ORDER BY id DESC"
        cursor.execute(query, params)
        data_tabel = cursor.fetchall()

        if filter_kategori:
            cursor.execute("SELECT * FROM users WHERE role != 'admin' AND kategori_magang = %s ORDER BY id DESC", (filter_kategori,))
        else:
            cursor.execute("SELECT * FROM users WHERE role != 'admin' ORDER BY id DESC")
        data_magang = cursor.fetchall()
        conn.close()

        # Rekap punya filter sendiri: rekap_cari, rekap_kategori, rekap_dari, rekap_sampai
        rekap_kategori = request.args.get('rekap_kategori', '')
        rekap_cari = request.args.get('rekap_cari', '').strip()
        rekap_dari, rekap_sampai = baca_rentang_rekap(request.args.get('rekap_dari', ''),
                                                      request.args.get('rekap_sampai', ''))
        rekap_bulanan, top_bulan = hitung_rekap_periode(rekap_dari, rekap_sampai, rekap_kategori, rekap_cari)
        top_hari_ini = top_tercepat_hari_ini(rekap_kategori, rekap_cari)
        daftar_rekap = ambil_daftar_rekap(rekap_dari, rekap_sampai, rekap_kategori, rekap_cari)

        ringkas = ringkasan_hari_ini(filter_kategori, filter_cari)

        total_absen = len(data_tabel)
        total_hadir = sum(1 for r in data_tabel if r['status'] in ['Masuk', 'Pulang', 'Hadir', 'Telat'])
        total_izin = sum(1 for r in data_tabel if r['status'] == 'Izin')
        total_sakit = sum(1 for r in data_tabel if r['status'] == 'Sakit')
        total_cuti = sum(1 for r in data_tabel if r['status'] == 'Cuti')
        total_tanpa_ket = sum(1 for r in data_tabel if r['status'] == 'Tanpa Keterangan')

        return render_template("admin_dashboard.html",
                               user=session,
                               ringkas=ringkas,
                               total_absen=total_absen,
                               total_hadir=total_hadir,
                               total_izin=total_izin,
                               total_sakit=total_sakit,
                               total_cuti=total_cuti,
                               total_tanpa_ket=total_tanpa_ket,
                               data_tabel=data_tabel,
                               data_magang=data_magang,
                               filter_tanggal=filter_tanggal,
                               filter_bulan=filter_bulan,
                               filter_tahun=filter_tahun,
                               filter_kategori=filter_kategori,
                               filter_cari=filter_cari,
                               rekap_kategori=rekap_kategori,
                               rekap_cari=rekap_cari,
                               rekap_dari=rekap_dari.strftime("%Y-%m-%d"),
                               rekap_sampai=rekap_sampai.strftime("%Y-%m-%d"),
                               rekap_label=label_periode(rekap_dari, rekap_sampai),
                               rekap_preset=[(n, d.strftime("%Y-%m-%d"), e.strftime("%Y-%m-%d")) for n, d, e in preset_rekap()],
                               rekap_bulanan=rekap_bulanan,
                               daftar_rekap=daftar_rekap,
                               top_bulan=top_bulan,
                               top_hari_ini=top_hari_ini)
    else:
        return render_user_dashboard()


@app.route("/admin/aksi/<int:id>", methods=["POST"])
def admin_aksi(id):
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    aksi = request.form.get("aksi")
    catatan = request.form.get("catatan_admin", "").strip() or "-"
    kembali = kembali_aman(request.form.get("kembali", ""))

    if aksi not in ("Disetujui", "Ditolak"):
        return redirect(kembali)

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        "UPDATE riwayat SET status_persetujuan = %s, catatan_admin = %s "
        "WHERE id = %s AND (status_persetujuan IS NULL OR status_persetujuan = 'Menunggu')",
        (aksi, catatan, id)
    )
    conn.commit()
    conn.close()
    return redirect(kembali)


# ---------------- ADMIN: SETUJUI SEMUA (mengikuti filter yang sedang aktif) ----------------
@app.route("/admin/setujui-semua", methods=["POST"])
def admin_setujui_semua():
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    kembali = kembali_aman(request.form.get("kembali", ""))

    filter_tanggal = request.form.get('tanggal', '')
    filter_bulan = request.form.get('bulan', '')
    filter_tahun = request.form.get('tahun', '')
    filter_kategori = request.form.get('kategori', '')
    filter_cari = request.form.get('cari', '').strip()

    query = ("UPDATE riwayat SET status_persetujuan = 'Disetujui' "
             "WHERE (status_persetujuan IS NULL OR status_persetujuan = '' OR status_persetujuan = 'Menunggu')")
    params = []
    query, params = tambah_filter_waktu(query, params, filter_tanggal, filter_bulan, filter_tahun)
    if filter_kategori:
        query += " AND kategori_magang = %s"
        params.append(filter_kategori)
    if filter_cari:
        query += " AND nama LIKE %s"
        params.append(f"%{filter_cari}%")

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(query, params)
    conn.commit()
    conn.close()
    return redirect(kembali)


# ---------------- ADMIN: HAPUS RIWAYAT SATU-SATU ----------------
@app.route("/admin/hapus/<int:id>", methods=["POST"])
def admin_hapus(id):
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    kembali = kembali_aman(request.form.get("kembali", ""))

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT bukti FROM riwayat WHERE id = %s", (id,))
    row = cursor.fetchone()
    if row:
        cursor.execute("DELETE FROM riwayat WHERE id = %s", (id,))
        conn.commit()
        hapus_file_upload(row.get('bukti'))
    conn.close()
    return redirect(kembali)


# ---------------- ADMIN: TAMBAH PESERTA ----------------
KATEGORI_VALID = ['MagangHub', 'PKL', 'Magang Berdampak', 'Job Familiarization', 'Magang Lainnya']


@app.route("/admin/peserta/tambah", methods=["POST"])
def admin_peserta_tambah():
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    kembali = kembali_aman(request.form.get("kembali", ""))

    username = request.form.get("username", "").strip()
    password = request.form.get("password", "").strip()
    nama = request.form.get("nama", "").strip()
    posisi = request.form.get("posisi", "").strip()
    asal_instansi = request.form.get("asal_instansi", "").strip()
    mentor = request.form.get("mentor", "").strip()
    email = request.form.get("email", "").strip()
    kategori = request.form.get("kategori_magang", "Magang Lainnya").strip()
    if kategori not in KATEGORI_VALID:
        kategori = 'Magang Lainnya'

    if not re.fullmatch(r"[A-Za-z0-9_.-]{3,50}", username):
        flash("Username 3-50 karakter, hanya huruf, angka, titik, garis bawah, atau strip.", "error")
        return redirect(kembali)
    if len(password) < 6:
        flash("Password minimal 6 karakter.", "error")
        return redirect(kembali)
    if not nama:
        flash("Nama lengkap wajib diisi.", "error")
        return redirect(kembali)

    # Foto wajah opsional. Tanpa foto, verifikasi wajah saat absen dilewati untuk peserta ini.
    nama_file_wajah = '-'
    path_wajah = None
    file_wajah = request.files.get("foto_wajah")
    if file_wajah and file_wajah.filename != '':
        aman = secure_filename(file_wajah.filename)
        if not aman.lower().endswith(EXT_GAMBAR):
            flash("Foto wajah harus berupa gambar (jpg, png, dan sejenisnya).", "error")
            return redirect(kembali)
        nama_file_wajah = f"reg_{username}_{aman}"
        path_wajah = os.path.join(app.config['UPLOAD_FOLDER'], nama_file_wajah)
        file_wajah.save(path_wajah)

    conn = None
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            "INSERT INTO users (username, password, nama, posisi, asal_instansi, mentor, email, role, foto_wajah, kategori_magang) "
            "VALUES (%s, %s, %s, %s, %s, %s, %s, 'user', %s, %s)",
            (username, password, nama, posisi, asal_instansi, mentor, email, nama_file_wajah, kategori)
        )
        conn.commit()
        flash(f"Peserta {nama} (username: {username}) berhasil ditambahkan.", "success")
    except pymysql.err.IntegrityError:
        if path_wajah and os.path.isfile(path_wajah):
            os.remove(path_wajah)
        flash("Username tersebut sudah terdaftar, gunakan username lain.", "error")
    except Exception as e:
        if path_wajah and os.path.isfile(path_wajah):
            os.remove(path_wajah)
        flash(f"Gagal menambahkan peserta: {e}", "error")
    finally:
        if conn:
            conn.close()
    return redirect(kembali)


# ---------------- ADMIN: HAPUS PESERTA ----------------
@app.route("/admin/peserta/hapus/<int:id>", methods=["POST"])
def admin_peserta_hapus(id):
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    kembali = kembali_aman(request.form.get("kembali", ""))

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT username, nama, role, foto_wajah FROM users WHERE id = %s", (id,))
    u = cursor.fetchone()
    if not u:
        flash("Peserta tidak ditemukan.", "error")
    elif u['role'] == 'admin':
        flash("Akun administrator tidak dapat dihapus.", "error")
    else:
        cursor.execute("DELETE FROM users WHERE id = %s", (id,))
        conn.commit()
        hapus_file_upload(u.get('foto_wajah'))   # hapus juga foto wajah pendaftaran
        flash(f"Peserta {u['nama']} (username: {u['username']}) dihapus. Riwayat absensinya tetap tersimpan.", "success")
    conn.close()
    return redirect(kembali)


# ---------------- ADMIN: EDIT PESERTA ----------------
@app.route("/admin/peserta/edit/<int:id>", methods=["POST"])
def admin_peserta_edit(id):
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    kembali = kembali_aman(request.form.get("kembali", ""))

    nama = request.form.get("nama", "").strip()
    posisi = request.form.get("posisi", "").strip()
    asal_instansi = request.form.get("asal_instansi", "").strip()
    mentor = request.form.get("mentor", "").strip()
    email = request.form.get("email", "").strip()
    kategori = request.form.get("kategori_magang", "Magang Lainnya").strip()
    password_baru = request.form.get("password", "").strip()
    if kategori not in KATEGORI_VALID:
        kategori = 'Magang Lainnya'

    if not nama:
        flash("Nama lengkap wajib diisi.", "error")
        return redirect(kembali)
    if password_baru and len(password_baru) < 6:
        flash("Password baru minimal 6 karakter.", "error")
        return redirect(kembali)

    path_baru = None
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM users WHERE id = %s", (id,))
        u = cursor.fetchone()
        if not u:
            flash("Peserta tidak ditemukan.", "error")
            return redirect(kembali)
        if u['role'] == 'admin':
            flash("Akun administrator tidak dapat diedit di sini.", "error")
            return redirect(kembali)

        # Riwayat absensi terhubung lewat nama, jadi nama baru tidak boleh sama dengan peserta lain
        if nama != u['nama']:
            cursor.execute("SELECT id FROM users WHERE nama = %s AND id != %s", (nama, id))
            if cursor.fetchone():
                flash("Nama tersebut sudah dipakai peserta lain.", "error")
                return redirect(kembali)

        # Foto wajah baru (opsional)
        foto_lama = u['foto_wajah']
        foto_baru = foto_lama
        file_wajah = request.files.get("foto_wajah")
        if file_wajah and file_wajah.filename != '':
            aman = secure_filename(file_wajah.filename)
            if not aman.lower().endswith(EXT_GAMBAR):
                flash("Foto wajah harus berupa gambar (jpg, png, dan sejenisnya).", "error")
                return redirect(kembali)
            foto_baru = f"reg_{u['username']}_{aman}"
            path_baru = os.path.join(app.config['UPLOAD_FOLDER'], foto_baru)
            file_wajah.save(path_baru)

        if password_baru:
            cursor.execute(
                "UPDATE users SET nama=%s, posisi=%s, asal_instansi=%s, mentor=%s, email=%s, "
                "kategori_magang=%s, foto_wajah=%s, password=%s WHERE id=%s",
                (nama, posisi, asal_instansi, mentor, email, kategori, foto_baru, password_baru, id))
        else:
            cursor.execute(
                "UPDATE users SET nama=%s, posisi=%s, asal_instansi=%s, mentor=%s, email=%s, "
                "kategori_magang=%s, foto_wajah=%s WHERE id=%s",
                (nama, posisi, asal_instansi, mentor, email, kategori, foto_baru, id))

        # Samakan riwayat lama supaya tetap terhubung dan rekap/filter kategori ikut benar
        cursor.execute(
            "UPDATE riwayat SET nama=%s, posisi=%s, kategori_magang=%s WHERE nama=%s",
            (nama, posisi, kategori, u['nama']))
        conn.commit()

        if foto_baru != foto_lama:
            hapus_file_upload(foto_lama)
        path_baru = None   # sukses, jangan dihapus di blok finally
        flash(f"Data peserta {nama} (username: {u['username']}) berhasil diperbarui.", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Gagal memperbarui peserta: {e}", "error")
    finally:
        if path_baru and os.path.isfile(path_baru):
            os.remove(path_baru)   # bersihkan foto baru jika proses gagal
        conn.close()
    return redirect(kembali)


# Segarkan data sesi dari database di setiap permintaan, sehingga perubahan dari admin
# (nama, kategori, dan lain-lain) langsung berlaku dan peserta yang dihapus otomatis keluar.
@app.before_request
def segarkan_sesi():
    if 'username' not in session or request.endpoint in (None, 'static', 'uploaded_file'):
        return
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM users WHERE username = %s", (session['username'],))
        u = cursor.fetchone()
        conn.close()
    except Exception:
        return
    if not u:
        session.clear()
        return redirect(url_for('login'))
    for key, val in u.items():
        session[key] = val


@app.route("/riwayat")
def riwayat_page():
    if 'username' not in session or session.get('role') != 'user':
        return redirect(url_for('login'))

    filter_tanggal = request.args.get('tanggal', '')
    filter_bulan = request.args.get('bulan', '')
    filter_tahun = request.args.get('tahun', '')

    conn = get_db_connection()
    cursor = conn.cursor()

    query = "SELECT * FROM riwayat WHERE nama = %s"
    params = [session['nama']]
    query, params = tambah_filter_waktu(query, params, filter_tanggal, filter_bulan, filter_tahun)

    query += " ORDER BY id DESC"
    cursor.execute(query, params)
    riwayat_user = cursor.fetchall()
    conn.close()

    return render_template("user_riwayat.html",
                           user=session,
                           riwayat=riwayat_user,
                           filter_tanggal=filter_tanggal,
                           filter_bulan=filter_bulan,
                           filter_tahun=filter_tahun)


# ---------------- PROFIL PESERTA (username & kategori magang bisa diubah) ----------------
@app.route("/profil", methods=["GET", "POST"])
def profil_page():
    if 'username' not in session or session.get('role') != 'user':
        return redirect(url_for('login'))

    if request.method == "GET":
        return render_template("user_profil.html", user=session)

    username_lama = session['username']
    nama_lama = session['nama']

    username = request.form.get("username", "").strip()
    nama = request.form.get("nama", "").strip()
    posisi = request.form.get("posisi", "").strip()
    asal_instansi = request.form.get("asal_instansi", "").strip()
    mentor = request.form.get("mentor", "").strip()
    email = request.form.get("email", "").strip()
    kategori = request.form.get("kategori_magang", "").strip()
    password = request.form.get("password", "").strip()
    file_wajah = request.files.get("foto_wajah")

    if kategori not in KATEGORI_VALID:
        kategori = session.get('kategori_magang') or 'Magang Lainnya'

    def balik(pesan):
        return render_template("user_profil.html", user=session, error=pesan)

    if not re.fullmatch(r"[A-Za-z0-9_.-]{3,50}", username):
        return balik("Username 3-50 karakter, hanya huruf, angka, titik, garis bawah, atau strip.")
    if not nama:
        return balik("Nama lengkap wajib diisi.")
    if password and len(password) < 6:
        return balik("Password baru minimal 6 karakter.")

    foto_lama = session.get('foto_wajah') or '-'
    foto_baru = foto_lama
    path_baru = None

    if file_wajah and file_wajah.filename != '':
        aman = secure_filename(file_wajah.filename)
        if not aman.lower().endswith(EXT_GAMBAR):
            return balik("Foto wajah harus berupa gambar (jpg, png, dan sejenisnya).")
        foto_baru = f"reg_{username}_{aman}"
        path_baru = os.path.join(app.config['UPLOAD_FOLDER'], foto_baru)

    conn = get_db_connection()
    sukses = False
    pesan_error = None
    try:
        cursor = conn.cursor()

        # Username harus unik
        if username != username_lama:
            cursor.execute("SELECT id FROM users WHERE username = %s", (username,))
            if cursor.fetchone():
                return balik("Username tersebut sudah dipakai, gunakan username lain.")

        # Riwayat absensi terhubung lewat nama, jadi nama tidak boleh sama dengan peserta lain
        if nama != nama_lama:
            cursor.execute("SELECT id FROM users WHERE nama = %s AND username != %s", (nama, username_lama))
            if cursor.fetchone():
                return balik("Nama tersebut sudah dipakai peserta lain.")

        if path_baru:
            file_wajah.save(path_baru)

        if password:
            cursor.execute(
                "UPDATE users SET username=%s, nama=%s, posisi=%s, asal_instansi=%s, mentor=%s, email=%s, "
                "kategori_magang=%s, foto_wajah=%s, password=%s WHERE username=%s",
                (username, nama, posisi, asal_instansi, mentor, email, kategori, foto_baru, password, username_lama))
        else:
            cursor.execute(
                "UPDATE users SET username=%s, nama=%s, posisi=%s, asal_instansi=%s, mentor=%s, email=%s, "
                "kategori_magang=%s, foto_wajah=%s WHERE username=%s",
                (username, nama, posisi, asal_instansi, mentor, email, kategori, foto_baru, username_lama))

        # Samakan riwayat lama supaya tetap terhubung dan rekap/filter kategori ikut benar
        cursor.execute(
            "UPDATE riwayat SET nama=%s, posisi=%s, kategori_magang=%s WHERE nama=%s",
            (nama, posisi, kategori, nama_lama))
        conn.commit()
        sukses = True
    except pymysql.err.IntegrityError:
        conn.rollback()
        pesan_error = "Username tersebut sudah dipakai, gunakan username lain."
    except Exception as e:
        conn.rollback()
        pesan_error = f"Gagal memperbarui profil: {e}"
    finally:
        conn.close()
        if not sukses and path_baru and os.path.isfile(path_baru):
            os.remove(path_baru)   # bersihkan foto baru jika proses gagal

    if not sukses:
        return balik(pesan_error)

    if foto_baru != foto_lama:
        hapus_file_upload(foto_lama)

    # Perbarui sesi supaya langsung berlaku
    session['username'] = username
    session['nama'] = nama
    session['posisi'] = posisi
    session['asal_instansi'] = asal_instansi
    session['mentor'] = mentor
    session['email'] = email
    session['kategori_magang'] = kategori
    session['foto_wajah'] = foto_baru
    return render_template("user_profil.html", user=session, success="Profil berhasil diperbarui!")


@app.route("/login", methods=["GET", "POST"])
def login():
    if 'username' in session:
        return redirect(url_for('index'))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        if len(password) < 6:
            return render_template("login.html", error="Password minimal 6 karakter.")
        try:
            conn = get_db_connection()
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM users WHERE username = %s AND password = %s", (username, password))
            user = cursor.fetchone()
            conn.close()

            if user:
                session.clear()
                for key, val in user.items():
                    session[key] = val
                return redirect(url_for('index'))
            else:
                return render_template("login.html", error="Username atau Password salah!")
        except Exception:
            return render_template("login.html", error="Kesalahan koneksi database.")
    return render_template("login.html")


# ---------------- LUPA PASSWORD (berdasarkan username) ----------------
@app.route("/lupa-password", methods=["GET", "POST"])
def lupa_password():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        pw = request.form.get("password", "").strip()
        pw2 = request.form.get("password2", "").strip()

        if not username:
            return render_template("lupa_password.html", error="Username wajib diisi.")
        if len(pw) < 6:
            return render_template("lupa_password.html", error="Password baru minimal 6 karakter.", username=username)
        if pw != pw2:
            return render_template("lupa_password.html", error="Konfirmasi password tidak sama.", username=username)

        try:
            conn = get_db_connection()
            cursor = conn.cursor()
            cursor.execute("SELECT id, role FROM users WHERE username = %s", (username,))
            user = cursor.fetchone()

            if not user:
                conn.close()
                return render_template("lupa_password.html", error="Username tidak terdaftar.", username=username)

            if user['role'] == 'admin':
                conn.close()
                return render_template("lupa_password.html",
                                       error="Password akun administrator tidak dapat direset di sini.",
                                       username=username)

            cursor.execute("UPDATE users SET password = %s WHERE id = %s", (pw, user['id']))
            conn.commit()
            conn.close()
            return render_template("lupa_password.html", success="Password berhasil diubah. Silakan login dengan password baru.")
        except Exception as e:
            print("[LUPA PASSWORD] Error:", repr(e))
            return render_template("lupa_password.html", error="Terjadi kesalahan koneksi database.", username=username)

    return render_template("lupa_password.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        nama = request.form.get("nama", "").strip()
        posisi = request.form.get("posisi", "").strip()
        asal_instansi = request.form.get("asal_instansi", "").strip()
        mentor = request.form.get("mentor", "").strip()
        email = request.form.get("email", "").strip()
        kategori_magang = request.form.get("kategori_magang", "Magang Lainnya").strip()

        if len(password) < 6:
            return render_template("register.html", error="Password minimal 6 karakter.")

        file_wajah = request.files.get("foto_wajah")
        if not file_wajah or file_wajah.filename == '':
            return render_template("register.html", error="Foto wajah registrasi wajib diunggah!")

        nama_file_wajah = f"reg_{username}_{file_wajah.filename}"
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], nama_file_wajah)
        file_wajah.save(filepath)

        conn = None
        try:
            conn = get_db_connection()
            cursor = conn.cursor()
            cursor.execute("SELECT COUNT(*) as cnt FROM users")
            res = cursor.fetchone()
            role = 'admin' if res['cnt'] == 0 else 'user'

            cursor.execute(
                "INSERT INTO users (username, password, nama, posisi, asal_instansi, mentor, email, role, foto_wajah, kategori_magang) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
                (username, password, nama, posisi, asal_instansi, mentor, email, role, nama_file_wajah, kategori_magang)
            )
            conn.commit()
            conn.close()
            return render_template("register.html", success="Registrasi berhasil! Silakan masuk.")
        except pymysql.err.IntegrityError:
            if conn:
                conn.close()
            return render_template("register.html", error="Username tersebut sudah terdaftar, silakan gunakan username lain.")
        except Exception as e:
            if conn:
                conn.close()
            return render_template("register.html", error=f"Terjadi kesalahan sistem: {str(e)}")

    return render_template("register.html")


def sudah_absen_hari_ini(nama, tanggal, daftar_status):
    """True jika peserta sudah punya catatan dengan salah satu status itu pada tanggal tsb."""
    conn = get_db_connection()
    cursor = conn.cursor()
    marker = ", ".join(["%s"] * len(daftar_status))
    cursor.execute(
        f"SELECT id FROM riwayat WHERE nama = %s AND waktu LIKE %s AND status IN ({marker}) LIMIT 1",
        [nama, tanggal + '%'] + list(daftar_status))
    ada = cursor.fetchone() is not None
    conn.close()
    return ada


@app.route("/absen", methods=["POST"])
def absen():
    if 'username' not in session or session.get('role') != 'user':
        return redirect(url_for('login'))

    nama = session['nama']
    posisi = session['posisi']
    kategori_magang = session.get('kategori_magang', 'Magang Lainnya')
    status = request.form.get("status")
    alasan = request.form.get("alasan", "-")

    waktu_sekarang_dt = sekarang()          # WITA
    waktu_sekarang_str = waktu_sekarang_dt.strftime("%Y-%m-%d %H:%M:%S")
    hari_ini = waktu_sekarang_dt.strftime("%Y-%m-%d")
    jam_sekarang = waktu_sekarang_dt.time()

    if status == 'Masuk':
        batas_masuk = datetime.strptime(BATAS_MASUK, "%H:%M:%S").time()
        if jam_sekarang > batas_masuk:
            status = 'Telat'

    elif status == 'Pulang':
        batas_pulang = datetime.strptime(BATAS_PULANG, "%H:%M:%S").time()
        if jam_sekarang < batas_pulang:
            return render_user_dashboard(
                error=f"Belum waktunya pulang! Absen pulang hanya dapat dilakukan setelah pukul {BATAS_PULANG[:5]} WITA."
            )

    bukti_simpan = "-"
    status_persetujuan = "Menunggu"
    user_lat, user_lon = KANTOR_LAT, KANTOR_LON

    if status in ['Masuk', 'Telat', 'Pulang']:
        # Absen datang/pulang hanya sekali per hari; besok otomatis kosong lagi
        if status in ('Masuk', 'Telat') and sudah_absen_hari_ini(nama, hari_ini, ('Masuk', 'Telat')):
            return render_user_dashboard(error="Anda sudah absen masuk hari ini.")
        if status == 'Pulang' and sudah_absen_hari_ini(nama, hari_ini, ('Pulang',)):
            return render_user_dashboard(error="Anda sudah absen pulang hari ini.")

        # Syarat utama: lokasi harus berada di dalam radius kantor
        try:
            user_lat = float(request.form.get("latitude", ""))
            user_lon = float(request.form.get("longitude", ""))
        except Exception:
            return render_user_dashboard(error="Lokasi GPS belum terbaca. Aktifkan GPS lalu coba lagi.")
        jarak = hitung_jarak_meter(user_lat, user_lon, KANTOR_LAT, KANTOR_LON)
        if jarak > MAX_RADIUS_METER:
            return render_user_dashboard(
                error=f"Anda berada {round(jarak)} m dari kantor. Absen hanya bisa dilakukan dalam radius {MAX_RADIUS_METER} m."
            )

        # Selfie (tanpa pencocokan wajah)
        img_data_base64 = request.form.get("foto_absen_base64", "")
        try:
            header, encoded = img_data_base64.split(",", 1)
            image_data = base64.b64decode(encoded)
            if not image_data:
                raise ValueError("kosong")
        except Exception:
            return render_user_dashboard(error="Foto selfie tidak terbaca. Silakan coba lagi.")

        nama_file_absen = f"absen_{status.lower()}_{waktu_sekarang_dt.strftime('%Y%m%d%H%M%S')}_{session['username']}.jpg"
        with open(os.path.join(app.config['UPLOAD_FOLDER'], nama_file_absen), "wb") as fh:
            fh.write(image_data)
        bukti_simpan = nama_file_absen
        metode_verifikasi = "Selfie + GPS (dalam radius kantor)"
    else:
        metode_verifikasi = "Pengajuan Digital"
        file_bukti = request.files.get("bukti")
        if file_bukti and file_bukti.filename != '':
            bukti_simpan = f"surat_{waktu_sekarang_dt.strftime('%Y%m%d%H%M%S')}_{secure_filename(file_bukti.filename)}"
            file_bukti.save(os.path.join(app.config['UPLOAD_FOLDER'], bukti_simpan))

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        """INSERT INTO riwayat (nama, posisi, status, alasan, bukti, waktu, latitude, longitude, metode_verifikasi, status_persetujuan, catatan_admin, kategori_magang)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
        (nama, posisi, status, alasan, bukti_simpan, waktu_sekarang_str, str(user_lat), str(user_lon), metode_verifikasi, status_persetujuan, "-", kategori_magang)
    )
    conn.commit()
    conn.close()
    return redirect(url_for('index'))


@app.route('/uploads/<filename>')
def uploaded_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)


@app.route("/export/excel")
def export_excel():
    """Export riwayat (admin: sesuai filter; user: hanya milik sendiri) lengkap dengan foto."""
    if 'username' not in session:
        return redirect(url_for('login'))

    filter_kategori = request.args.get('kategori', '')
    filter_cari = request.args.get('cari', '')
    filter_tanggal = request.args.get('tanggal', '')
    filter_bulan = request.args.get('bulan', '')
    filter_tahun = request.args.get('tahun', '')

    conn = get_db_connection()
    cursor = conn.cursor()

    query = "SELECT * FROM riwayat WHERE 1=1"
    params = []

    if session.get('role') == 'admin':
        if filter_kategori:
            query += " AND kategori_magang = %s"
            params.append(filter_kategori)
        if filter_cari:
            query += " AND nama LIKE %s"
            params.append(f"%{filter_cari}%")
    else:
        query += " AND nama = %s"
        params.append(session.get('nama'))

    query, params = tambah_filter_waktu(query, params, filter_tanggal, filter_bulan, filter_tahun)
    query += " ORDER BY waktu DESC"

    cursor.execute(query, params)
    data = cursor.fetchall()
    conn.close()

    wb = Workbook()
    wb.remove(wb.active)
    tulis_sheet_daftar(wb, data, "Riwayat Absensi")

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    if session.get('role') == 'admin':
        suffix = filter_kategori.replace(' ', '_') if filter_kategori else 'semua'
    else:
        suffix = 'pribadi'
    nama_file = f"Riwayat_Absensi_{suffix}_{sekarang().strftime('%Y%m%d_%H%M')}.xlsx"

    return output.getvalue(), 200, {
        "Content-Disposition": f"attachment; filename={nama_file}",
        "Content-Type": XLSX_MIME
    }

@app.route("/export/rekap")
def export_rekap():
    """Export rekap absensi (sesuai filter) + daftar lengkap dengan foto. Khusus admin."""
    if 'username' not in session or session.get('role') != 'admin':
        return redirect(url_for('login'))

    kategori = request.args.get('rekap_kategori', '')
    cari = request.args.get('rekap_cari', '').strip()
    dari, sampai = baca_rentang_rekap(request.args.get('rekap_dari', ''),
                                      request.args.get('rekap_sampai', ''))
    rekap, _ = hitung_rekap_periode(dari, sampai, kategori, cari)
    daftar = ambil_daftar_rekap(dari, sampai, kategori, cari)

    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Absensi"

    ws.append(["REKAP ABSENSI MAGANG - KSOP UTAMA MAKASSAR"])
    ws.append([f"Periode: {label_periode(dari, sampai)}   |   Kategori: {kategori or 'Semua Kategori'}   |   Nama: {cari or 'Semua'}"])
    ws.append([])
    header = ["No", "Nama", "Kategori Magang", "Hadir (hari)", "Telat (hari)", "Izin", "Sakit",
              "Cuti", "Tanpa Keterangan", "Rata-rata Jam Masuk", "Lencana", "Rincian Tanggal & Jam"]
    ws.append(header)
    ws['A1'].font = Font(bold=True, size=13, color="0B2545")
    ws['A2'].font = Font(italic=True, color="64748B")
    for col in range(1, len(header) + 1):
        c = ws.cell(row=4, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="0B2545", end_color="0B2545", fill_type="solid")
        c.alignment = Alignment(horizontal="center")

    if not rekap:
        ws.append(["", "Tidak ada peserta pada filter yang dipilih"])
    for i, r in enumerate(rekap, start=1):
        rincian_txt = "\n".join(
            f"{h['tanggal']}  masuk {h['masuk']}  pulang {h['pulang']}  ({h['ket']})" for h in r['rincian']
        ) or "-"
        ws.append([i, r['nama'], r['kategori'], r['hadir'], r['telat'], r['izin'], r['sakit'],
                   r['cuti'], r['tanpa_ket'], r['rata_jam'],
                   f"Peringkat {r['lencana']}" if r['lencana'] else "-", rincian_txt])
        ws.cell(row=ws.max_row, column=12).alignment = Alignment(wrap_text=True, vertical="top")

    for i, w in enumerate([6, 26, 20, 13, 13, 8, 8, 8, 18, 20, 14, 58], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Sheet 2: semua absensi sesuai filter + foto
    tulis_sheet_daftar(wb, daftar, "Daftar Absensi")

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    suffix = kategori.replace(' ', '_') if kategori else 'semua'
    nama_file = f"Rekap_Absensi_{suffix}_{dari.strftime('%Y%m%d')}_{sampai.strftime('%Y%m%d')}.xlsx"
    return output.getvalue(), 200, {
        "Content-Disposition": f"attachment; filename={nama_file}",
        "Content-Type": XLSX_MIME
    }


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for('login'))


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)